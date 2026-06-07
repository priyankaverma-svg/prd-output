#!/usr/bin/env python3
"""Generate BRD docx from markdown. Main body hard cap: 6 pages. Appendix excluded."""

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor

MAX_BODY_PAGES = 6
FONT_NAME = "Calibri"
FONT_SIZE = Pt(10)
HEADING1_SIZE = Pt(14)
HEADING2_SIZE = Pt(11)
MARGIN = Inches(0.75)
APPENDIX_MARKERS = ("## Appendix", "## appendix")


def setup_document():
    doc = Document()
    for section in doc.sections:
        section.top_margin = MARGIN
        section.bottom_margin = MARGIN
        section.left_margin = MARGIN
        section.right_margin = MARGIN
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = FONT_SIZE
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = 1.0
    return doc


def split_body_appendix(md_text):
    for marker in APPENDIX_MARKERS:
        idx = md_text.find(marker)
        if idx != -1:
            return md_text[:idx].rstrip(), md_text[idx:].strip()
    return md_text, ""


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = FONT_NAME
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.size = HEADING1_SIZE if level == 1 else HEADING2_SIZE
    h.paragraph_format.space_before = Pt(6)
    h.paragraph_format.space_after = Pt(3)
    return h


def add_paragraph(doc, text, bold=False, bullet=False):
    if bullet:
        p = doc.add_paragraph(text, style="List Bullet")
    else:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        run.font.name = FONT_NAME
        run.font.size = FONT_SIZE
    p.paragraph_format.space_after = Pt(2)
    return p


def parse_table_block(lines):
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.split("|")[1:-1]]
        if cells and not all(set(c) <= set("-:") for c in cells):
            rows.append(cells)
    return rows


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            if j < cols:
                table.rows[i].cells[j].text = cell
                for p in table.rows[i].cells[j].paragraphs:
                    for run in p.runs:
                        run.font.name = FONT_NAME
                        run.font.size = Pt(9)
    doc.add_paragraph()


def estimate_pages(doc):
    chars = sum(len(p.text) for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                chars += len(cell.text)
    return max(1, round(chars / 3000 + 0.5))


def parse_markdown(md_text):
    lines = md_text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            blocks.append(("h1", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
        elif line.startswith("### "):
            blocks.append(("h3", line[4:].strip()))
        elif line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            blocks.append(("table", parse_table_block(table_lines)))
            continue
        elif line.startswith("- ") or line.startswith("* "):
            blocks.append(("bullet", line[2:].strip()))
        elif re.match(r"^\d+\.\s", line):
            blocks.append(("numbered", line.strip()))
        elif line.strip() == "---":
            pass
        elif line.strip():
            blocks.append(("text", line.strip()))
        i += 1
    return blocks


def build_doc(blocks, page_break_before=False):
    doc = setup_document()
    if page_break_before and blocks:
        p = doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
    for kind, content in blocks:
        if kind == "h1":
            add_heading(doc, content, level=1)
        elif kind == "h2":
            add_heading(doc, content, level=2)
        elif kind == "h3":
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.bold = True
            run.font.name = FONT_NAME
            run.font.size = FONT_SIZE
        elif kind == "table":
            add_table(doc, content)
        elif kind == "bullet":
            add_paragraph(doc, content, bullet=True)
        elif kind == "numbered":
            add_paragraph(doc, content, bullet=True)
        elif kind == "text":
            add_paragraph(doc, content)
    return doc


def merge_docs(body_doc, appendix_doc):
    for element in appendix_doc.element.body:
        body_doc.element.body.append(element)
    return body_doc


def main():
    parser = argparse.ArgumentParser(description="Generate BRD docx — 6-page body + appendix")
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--max-body-pages", type=int, default=MAX_BODY_PAGES)
    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)

    if not input_path.exists():
        print(f"ERROR: Input file not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    md_text = input_path.read_text(encoding="utf-8")
    body_md, appendix_md = split_body_appendix(md_text)

    body_blocks = parse_markdown(body_md)
    body_doc = build_doc(body_blocks)
    body_pages = estimate_pages(body_doc)

    if appendix_md:
        appendix_blocks = parse_markdown(appendix_md)
        appendix_doc = build_doc(appendix_blocks, page_break_before=True)
        doc = merge_docs(body_doc, appendix_doc)
        appendix_pages = estimate_pages(appendix_doc)
    else:
        doc = body_doc
        appendix_pages = 0

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))

    print(f"Saved: {output_path}")
    print(f"Estimated main body pages: {body_pages} (limit: {args.max_body_pages})")
    if appendix_pages:
        print(f"Estimated appendix pages: {appendix_pages} (not counted toward limit)")

    if body_pages > args.max_body_pages:
        print(
            f"WARNING: Main body may exceed {args.max_body_pages}-page limit. "
            "Compress: FAQs → financial prose → market narrative → evidence. Re-run after edits.",
            file=sys.stderr,
        )
        sys.exit(2)

    sys.exit(0)


if __name__ == "__main__":
    main()
